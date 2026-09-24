"""Write parsed statements to an Excel workbook, one sheet per statement plus a summary."""

import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Indian digit grouping: 1,23,45,678.00
INR_FORMAT = '[>=10000000]##\\,##\\,##\\,##0.00;[>=100000]##\\,##\\,##0.00;##,##0.00'
DATE_FORMAT = "DD-MM-YYYY"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
BAD_FILL = PatternFill("solid", fgColor="F8CBAD")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _money_format(value):
    # Indian grouping formats only cover positives; overdrawn balances use a plain format
    return INR_FORMAT if value is None or value >= 0 else "#,##0.00"


def _sheet_name(stmt, used):
    base = stmt.account_no or re.sub(r"\.pdf$", "", stmt.source.split("/")[-1], flags=re.I)
    base = re.sub(r"[\\/*?:\[\]]", "_", base)[:28] or "Statement"
    name, n = base, 2
    while name.lower() in used:
        name = f"{base[:25]}_{n}"
        n += 1
    used.add(name.lower())
    return name


def _header(ws, row, titles):
    for c, title in enumerate(titles, 1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _write_transactions(ws, stmt):
    txns = stmt.transactions
    show_value = any(t.value_date for t in txns)
    show_ref = any(t.ref for t in txns)
    cols = [("Date", 12)]
    if show_value:
        cols.append(("Value Date", 12))
    cols.append(("Description / Narration", 60))
    if show_ref:
        cols.append(("Chq. / Ref. No.", 22))
    cols += [("Debit (Dr)", 16), ("Credit (Cr)", 16), ("Balance", 18)]
    cols += [(name, 14) for name in stmt.extra_columns]
    cols.append(("Balance Check", 14))
    idx = {name: i + 1 for i, (name, _) in enumerate(cols)}

    _header(ws, 1, [c[0] for c in cols])
    for i, (_, width) in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 30

    r = 2
    ws.cell(row=r, column=idx["Description / Narration"], value="Opening Balance").font = Font(bold=True)
    if stmt.opening_balance is not None:
        c = ws.cell(row=r, column=idx["Balance"], value=stmt.opening_balance)
        c.number_format, c.font = _money_format(stmt.opening_balance), Font(bold=True)
    first_txn_row = r + 1

    for t in txns:
        r += 1
        values = {
            "Date": t.date, "Value Date": t.value_date, "Description / Narration": t.description,
            "Chq. / Ref. No.": t.ref, "Debit (Dr)": t.debit, "Credit (Cr)": t.credit,
            "Balance": t.balance, "Balance Check": t.check,
        }
        values.update(t.extra)
        for name, col in idx.items():
            v = values.get(name)
            if v in ("", None):
                continue
            cell = ws.cell(row=r, column=col, value=v)
            if name in ("Date", "Value Date"):
                cell.number_format = DATE_FORMAT
            elif name in ("Debit (Dr)", "Credit (Cr)", "Balance"):
                cell.number_format = _money_format(v)
        ws.cell(row=r, column=idx["Description / Narration"]).alignment = Alignment(wrap_text=True, vertical="top")
        if t.check == "Mismatch":
            for col in range(1, len(cols) + 1):
                ws.cell(row=r, column=col).fill = BAD_FILL

    last_txn_row = r
    r += 1
    ws.cell(row=r, column=idx["Description / Narration"], value="Total")
    for name in ("Debit (Dr)", "Credit (Cr)"):
        L = get_column_letter(idx[name])
        cell = ws.cell(row=r, column=idx[name], value=f"=SUM({L}{first_txn_row}:{L}{last_txn_row})")
        cell.number_format = INR_FORMAT
    if stmt.closing_balance is not None:
        c = ws.cell(row=r, column=idx["Balance"], value=stmt.closing_balance)
        c.number_format = _money_format(stmt.closing_balance)
    for col in range(1, len(cols) + 1):
        cell = ws.cell(row=r, column=col)
        cell.fill, cell.font = TOTAL_FILL, Font(bold=True)

    for row in ws.iter_rows(min_row=2, max_row=r, max_col=len(cols)):
        for cell in row:
            cell.border = BORDER
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last_txn_row}"


def write_workbook(statements, path):
    """Save statements (list of Statement) to `path`."""
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    titles = ["Sheet", "File", "Bank", "Account No.", "Period", "Opening Balance", "Total Debit",
              "Total Credit", "Closing Balance", "Computed Closing", "Transactions",
              "Rows to review", "Notes"]
    _header(summary, 1, titles)
    widths = [18, 30, 22, 18, 26, 18, 18, 18, 18, 18, 13, 13, 70]
    for i, w in enumerate(widths, 1):
        summary.column_dimensions[get_column_letter(i)].width = w
    summary.row_dimensions[1].height = 30

    used = {"summary"}
    for n, stmt in enumerate(statements, 2):
        name = _sheet_name(stmt, used)
        _write_transactions(wb.create_sheet(name), stmt)
        computed = None
        if stmt.opening_balance is not None:
            computed = round(stmt.opening_balance - stmt.total_debit + stmt.total_credit, 2)
        row = [name, stmt.source.split("/")[-1], stmt.bank, stmt.account_no, stmt.period,
               stmt.opening_balance, stmt.total_debit, stmt.total_credit, stmt.closing_balance,
               computed, len(stmt.transactions), stmt.mismatches, " ".join(stmt.warnings)]
        for c, v in enumerate(row, 1):
            cell = summary.cell(row=n, column=c, value=v)
            cell.border = BORDER
            if 6 <= c <= 10 and v is not None:
                cell.number_format = _money_format(v)
        summary.cell(row=n, column=1).hyperlink = f"#'{name}'!A1"
        summary.cell(row=n, column=1).font = Font(color="0563C1", underline="single")
        summary.cell(row=n, column=13).alignment = Alignment(wrap_text=True, vertical="top")
        if stmt.mismatches or (computed is not None and stmt.closing_balance is not None
                               and abs(computed - stmt.closing_balance) > 0.005):
            summary.cell(row=n, column=12).fill = BAD_FILL
    summary.freeze_panes = "A2"
    wb.save(path)
