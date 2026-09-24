#!/usr/bin/env python3
"""Export starred Gmail messages to Excel.

Columns: Date | Time | Subject | Link | Reg No

Reg No is the vehicle registration number(s) found in the subject line
(blank if none; several are comma-separated).

Usage:
    python gmail_starred_export.py                       # -> starred_mails.xlsx
    python gmail_starred_export.py -o out.xlsx --tz Asia/Kolkata
"""

import argparse
import os
import sys
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from regno import find_reg_numbers

SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]
HEADERS = ["Date", "Time", "Subject", "Link", "Reg No"]


def gmail_service(credentials_file, token_file):
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    creds = None
    if os.path.exists(token_file):
        creds = Credentials.from_authorized_user_file(token_file, SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not os.path.exists(credentials_file):
                sys.exit(
                    f"Missing {credentials_file}. Download an OAuth 'Desktop app' "
                    "client from Google Cloud Console (see README.md)."
                )
            flow = InstalledAppFlow.from_client_secrets_file(credentials_file, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(token_file, "w") as f:
            f.write(creds.to_json())
    return build("gmail", "v1", credentials=creds, cache_discovery=False)


def fetch_starred(service):
    """Yield (message_id, thread_id, internal_date_ms, subject) for every starred message."""
    page_token = None
    while True:
        resp = (
            service.users()
            .messages()
            .list(userId="me", labelIds=["STARRED"], maxResults=500, pageToken=page_token)
            .execute()
        )
        for ref in resp.get("messages", []):
            msg = (
                service.users()
                .messages()
                .get(userId="me", id=ref["id"], format="metadata", metadataHeaders=["Subject"])
                .execute()
            )
            headers = msg.get("payload", {}).get("headers", [])
            subject = next((h["value"] for h in headers if h["name"].lower() == "subject"), "")
            yield msg["id"], msg["threadId"], int(msg["internalDate"]), subject
        page_token = resp.get("nextPageToken")
        if not page_token:
            break


def mail_link(email, msg_id, thread_id):
    """Same link format Gmail uses; ``authuser`` opens the right account even
    when several are signed in."""
    return (
        f"https://mail.google.com/mail/?authuser={email}"
        f"#all/thread-f:{int(thread_id, 16)}|msg-f:{int(msg_id, 16)}"
    )


def build_rows(messages, tz, email):
    rows = []
    # newest first
    for msg_id, thread_id, internal_ms, subject in sorted(messages, key=lambda m: m[2], reverse=True):
        dt = datetime.fromtimestamp(internal_ms / 1000, tz=timezone.utc).astimezone(tz)
        link = mail_link(email, msg_id, thread_id)
        rows.append(
            [
                dt.strftime("%d-%m-%Y"),
                dt.strftime("%H:%M:%S"),
                subject,
                link,
                ", ".join(find_reg_numbers(subject)),
            ]
        )
    return rows


def write_excel(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Starred Mails"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")

    for row in rows:
        ws.append(row)
        link_cell = ws.cell(row=ws.max_row, column=4)
        link_cell.hyperlink = row[3]
        link_cell.style = "Hyperlink"

    for i, width in enumerate([12, 10, 70, 55, 20], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("-o", "--output", default="starred_mails.xlsx")
    p.add_argument("--tz", default="Asia/Kolkata", help="Timezone for Date/Time (default: Asia/Kolkata)")
    p.add_argument("--credentials", default="credentials.json")
    p.add_argument("--token", default="token.json")
    args = p.parse_args()

    service = gmail_service(args.credentials, args.token)
    email = service.users().getProfile(userId="me").execute()["emailAddress"]
    rows = build_rows(fetch_starred(service), ZoneInfo(args.tz), email)
    write_excel(rows, args.output)
    with_reg = sum(1 for r in rows if r[4])
    print(f"Saved {len(rows)} starred mails to {args.output} ({with_reg} with a Reg No).")


if __name__ == "__main__":
    main()
