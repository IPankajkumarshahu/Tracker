"""Tests for the bank statement PDF -> Excel converter.

The statements are drawn with reportlab (pip install reportlab) the way banks lay
them out: text placed at fixed positions, wrapped narrations, right-aligned amounts.
"""

import datetime as dt
import os
import random
import tempfile
import unittest

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from bank_statement import StatementError, parse_statement, write_workbook
from bank_statement.parser import parse_amount, parse_date

FONT, SIZE = "Helvetica", 8

NARRATIONS = [
    "UPI/DR/412345678901/RAHUL SHARMA/HDFC/rahul.sharma@okhdfcbank/Payment from PhonePe",
    "NEFT CR-SBIN0001234-ACME TECHNOLOGIES PRIVATE LIMITED-SALARY FOR APRIL 2024",
    "ATW-512345XXXXXX1234-S1ANMU12-MUMBAI",
    "IMPS-412398765432-GUPTA TRADERS-KKBK-XXXXXXXX1234-INVOICE 245",
    "POS 512345XXXXXX1234 AMAZON PAY INDIA PVT LTD",
    "CHQ DEP - CLG - 000123 - MEHTA AND ASSOCIATES CHARTERED ACCOUNTANTS",
    "ACH D- TP ACH ICICI PRUDENTIAL-1234567890",
    "INTEREST PAID TILL 30-JUN-2024",
    "GST PAYMENT CBDT 0510012 CIN 24063000123456",
    "RTGS CR-HDFC0000001-RELIANCE RETAIL LTD-UTR HDFCR52024061112345678 AGAINST BILL NO 7781",
    "SMS CHARGES",
    "BIL/ONL/000987654/ELECTRICITY BOARD/MSEDCL",
]


def make_transactions(n, opening=85000.0, start=dt.date(2024, 4, 1), seed=1):
    rnd = random.Random(seed)
    bal, day, txns = opening, start, []
    for i in range(n):
        day += dt.timedelta(days=rnd.choice([0, 1, 1, 2]))
        amt = round(rnd.choice([rnd.uniform(10, 999), rnd.uniform(1000, 99999), rnd.uniform(100000, 450000)]), 2)
        credit = rnd.random() < 0.4 or bal - amt < 0
        bal = round(bal + amt if credit else bal - amt, 2)
        txns.append({
            "date": day, "value_date": day, "desc": NARRATIONS[i % len(NARRATIONS)],
            "ref": f"{rnd.randint(10**11, 10**12 - 1)}" if i % 3 else "",
            "debit": None if credit else amt, "credit": amt if credit else None, "balance": bal,
        })
    return txns


def money(v):
    return f"{v:,.2f}"


def wrap(text, width, txn=None):
    """Word-wrap like a bank does; over-long words are cut. Records the expected
    narration (wrapped lines joined by a space) on `txn`."""
    words = []
    for word in text.split(" "):
        while stringWidth(word, FONT, SIZE) > width:
            n = len(word)
            while stringWidth(word[:n], FONT, SIZE) > width:
                n -= 1
            words.append(word[:n])
            word = word[n:]
        words.append(word)
    lines, cur = [], ""
    for word in words:
        trial = (cur + " " + word).strip()
        if stringWidth(trial, FONT, SIZE) <= width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = word
    lines.append(cur)
    if txn is not None:
        txn["shown"] = " ".join(lines)
    return lines


class Pdf:
    """Minimal drawing helper: columns are (key, label, x0, x1, align)."""

    def __init__(self, path, pagesize=A4, password=None):
        enc = StandardEncryption(password, canPrint=1) if password else None
        self.c = canvas.Canvas(path, pagesize=pagesize, encrypt=enc)
        self.w, self.h = pagesize
        self.c.setFont(FONT, SIZE)

    def text(self, x, y, s, align="left", bold=False):
        self.c.setFont(FONT + ("-Bold" if bold else ""), SIZE)
        if align == "right":
            self.c.drawRightString(x, self.h - y, s)
        else:
            self.c.drawString(x, self.h - y, s)
        self.c.setFont(FONT, SIZE)

    def hline(self, x0, x1, y):
        self.c.line(x0, self.h - y, x1, self.h - y)

    def page(self):
        self.c.showPage()
        self.c.setFont(FONT, SIZE)

    def save(self):
        self.c.save()


# --------------------------------------------------------------- layouts

def hdfc_like(path, txns, opening, password=None):
    """No ruling lines, dd/mm/yy dates, top-aligned wrapped narration, header on each page."""
    pdf = Pdf(path, pagesize=landscape(A4), password=password)
    cols = [("date", "Date", 30, 75, "left"), ("desc", "Narration", 80, 330, "left"),
            ("ref", "Chq./Ref.No.", 340, 420, "left"), ("value_date", "Value Dt", 430, 475, "left"),
            ("debit", "Withdrawal Amt.", 480, 560, "right"), ("credit", "Deposit Amt.", 570, 650, "right"),
            ("balance", "Closing Balance", 660, 750, "right")]
    lh = 10

    def header():
        pdf.text(30, 30, "HDFC BANK Ltd.", bold=True)
        pdf.text(30, 42, "Account No : 50100123456789   IFSC : HDFC0001234")
        pdf.text(30, 54, "Statement From : 01/04/2024 To : 30/06/2024")
        for _, label, x0, x1, align in cols:
            pdf.text(x1 if align == "right" else x0, 80, label, align, bold=True)
        return 96

    y, pno = header(), 1
    for t in txns:
        lines = wrap(t["desc"], 245, t)
        if y + lh * len(lines) > 560:
            pdf.text(380, 580, f"Page No .: {pno}")
            pdf.page()
            pno += 1
            y = header()
        pdf.text(30, y, t["date"].strftime("%d/%m/%y"))
        pdf.text(340, y, t["ref"].zfill(16) if t["ref"] else "0000000000000000")
        pdf.text(430, y, t["value_date"].strftime("%d/%m/%y"))
        for key in ("debit", "credit", "balance"):
            if t[key] is not None:
                x1 = next(c[3] for c in cols if c[0] == key)
                pdf.text(x1, y, money(t[key]), "right")
        for i, ln in enumerate(lines):
            pdf.text(80, y + i * lh, ln)
        y += lh * len(lines)
    pdf.text(30, y + 30, "STATEMENT SUMMARY :-", bold=True)
    pdf.text(30, y + 42, "Opening Balance   Dr Count   Cr Count   Debits   Credits   Closing Bal")
    pdf.text(30, y + 54, f"{money(opening)}   5   3   1,000.00   2,000.00   {money(txns[-1]['balance'])}")
    pdf.text(380, 580, f"Page No .: {pno}")
    pdf.save()


def sbi_like(path, txns, opening):
    """Ruled rows, 'd Mon yyyy' dates, two-line headers, opening balance row."""
    pdf = Pdf(path)
    cols = [("date", ["Txn Date"], 25, 75, "left"), ("value_date", ["Value", "Date"], 80, 130, "left"),
            ("desc", ["Description"], 135, 300, "left"), ("ref", ["Ref No./Cheque", "No."], 305, 385, "left"),
            ("debit", ["Debit"], 390, 450, "right"), ("credit", ["Credit"], 455, 515, "right"),
            ("balance", ["Balance"], 520, 580, "right")]
    lh = 10

    def header(y):
        pdf.hline(22, 583, y - 10)
        for _, labels, x0, x1, align in cols:
            for i, lab in enumerate(labels):
                pdf.text(x1 if align == "right" else x0, y + i * lh, lab, align, bold=True)
        pdf.hline(22, 583, y + 16)
        return y + 28

    pdf.text(25, 30, "State Bank of India", bold=True)
    pdf.text(25, 42, "Account Number : 00000031234567890   IFSC Code : SBIN0004321")
    pdf.text(25, 54, "Account Statement from 1 Apr 2024 to 30 Jun 2024")
    y = header(90)
    pdf.text(135, y, "Opening Balance")
    pdf.text(580, y, money(opening), "right")
    y += 12
    pdf.hline(22, 583, y - 8)
    for t in txns:
        lines = wrap(t["desc"], 160, t)
        ref_lines = wrap(t["ref"], 75) if t["ref"] else []
        rows = max(len(lines), len(ref_lines), 1)
        if y + rows * lh > 800:
            pdf.page()
            y = header(40)
        pdf.text(25, y, t["date"].strftime("%-d %b %Y"))
        pdf.text(80, y, t["value_date"].strftime("%-d %b %Y"))
        for i, ln in enumerate(lines):
            pdf.text(135, y + i * lh, ln)
        for i, ln in enumerate(ref_lines):
            pdf.text(305, y + i * lh, ln)
        for key in ("debit", "credit", "balance"):
            if t[key] is not None:
                x1 = next(c[3] for c in cols if c[0] == key)
                pdf.text(x1, y, money(t[key]), "right")
        y += rows * lh + 4
        pdf.hline(22, 583, y - 8)
    pdf.text(25, y + 20, "**This is a computer generated statement and does not require a signature")
    pdf.save()


def kotak_like(path, txns, opening):
    """Single Amount column with (Dr)/(Cr), balance with (Cr), dates centred in the row,
    a Sl. No. column, no ruling lines but extra space between rows."""
    pdf = Pdf(path)
    cols = [("sl", "Sl. No.", 25, 50, "left"), ("date", "Date", 55, 105, "left"),
            ("desc", "Narration", 110, 330, "left"), ("ref", "Chq/Ref No", 335, 420, "left"),
            ("amount", "Amount", 425, 500, "right"), ("balance", "Balance", 505, 580, "right")]
    lh = 10

    def header(y):
        for _, label, x0, x1, align in cols:
            pdf.text(x1 if align == "right" else x0, y, label, align, bold=True)
        return y + 18

    pdf.text(25, 30, "Kotak Mahindra Bank", bold=True)
    pdf.text(25, 42, "Account No. 1234567890   IFSC KKBK0000958")
    pdf.text(25, 54, "Period 01-04-2024 to 30-06-2024")
    y = header(80)
    for n, t in enumerate(txns, 1):
        lines = wrap(t["desc"], 215, t)
        block = len(lines) * lh
        if y + block > 800:
            pdf.page()
            y = header(40)
        mid = y + (len(lines) - 1) * lh / 2
        pdf.text(25, mid, str(n))
        pdf.text(55, mid, t["date"].strftime("%d-%m-%Y"))
        for i, ln in enumerate(lines):
            pdf.text(110, y + i * lh, ln)
        pdf.text(335, mid, t["ref"])
        amt = f"{money(t['debit'])}(Dr)" if t["debit"] else f"{money(t['credit'])}(Cr)"
        pdf.text(500, mid, amt, "right")
        pdf.text(580, mid, f"{money(t['balance'])}(Cr)", "right")
        y += block + 8
    pdf.save()


def icici_like(path, txns, opening, header_every_page=False):
    """Newest first, dd-Mon-yyyy, '(INR )' header suffixes, header on first page only."""
    pdf = Pdf(path)
    cols = [("sno", ["S No."], 25, 50, "left"), ("date", ["Transaction", "Date"], 55, 110, "left"),
            ("desc", ["Transaction Remarks"], 115, 340, "left"),
            ("debit", ["Withdrawal", "Amount (INR )"], 345, 425, "right"),
            ("credit", ["Deposit", "Amount (INR )"], 430, 505, "right"),
            ("balance", ["Balance (INR )"], 510, 580, "right")]
    lh = 10

    def header(y):
        for _, labels, x0, x1, align in cols:
            for i, lab in enumerate(labels):
                pdf.text(x1 if align == "right" else x0, y + i * lh, lab, align, bold=True)
        return y + 26

    pdf.text(25, 30, "ICICI Bank Limited", bold=True)
    pdf.text(25, 42, "Account Number 000401234567   IFSC ICIC0000004")
    y = header(70)
    for n, t in enumerate(reversed(txns), 1):
        lines = wrap(t["desc"], 220, t)
        if y + len(lines) * lh > 800:
            pdf.page()
            y = header(40) if header_every_page else 40
        pdf.text(25, y, str(n))
        pdf.text(55, y, t["date"].strftime("%d-%b-%Y"))
        for i, ln in enumerate(lines):
            pdf.text(115, y + i * lh, ln)
        for key in ("debit", "credit", "balance"):
            x1 = next(c[3] for c in cols if c[0] == key)
            pdf.text(x1, y, money(t[key]) if t[key] is not None else "0.00", "right")
        y += len(lines) * lh + 3
    pdf.save()


def drcr_column_like(path, txns, opening, with_flag=True):
    """Amount + separate Dr/Cr column (or no flag at all), plus an extra 'Init. Br' column."""
    pdf = Pdf(path)
    cols = [("date", "Tran Date", 25, 75, "left"), ("ref", "Chq No", 80, 140, "left"),
            ("desc", "Particulars", 145, 360, "left"), ("amount", "Amount", 365, 440, "right")]
    if with_flag:
        cols.append(("flag", "Dr/Cr", 450, 475, "left"))
    cols += [("balance", "Balance", 480, 550, "right"), ("br", "Init. Br", 555, 585, "left")]
    lh = 10
    pdf.text(25, 30, "AXIS BANK", bold=True)
    pdf.text(25, 42, "Customer ID : 123456789  A/C No: 917010012345678  IFSC: UTIB0000123")
    for _, label, x0, x1, align in cols:
        pdf.text(x1 if align == "right" else x0, 70, label, align, bold=True)
    y = 86
    pdf.text(145, y, "OPENING BALANCE")
    pdf.text(550, y, money(opening), "right")
    y += lh
    for t in txns:
        lines = wrap(t["desc"], 210, t)
        pdf.text(25, y, t["date"].strftime("%d-%m-%Y"))
        pdf.text(80, y, t["ref"][:6])
        for i, ln in enumerate(lines):
            pdf.text(145, y + i * lh, ln)
        amt = t["debit"] or t["credit"]
        pdf.text(440, y, money(amt), "right")
        if with_flag:
            pdf.text(450, y, "DR" if t["debit"] else "CR")
        pdf.text(550, y, money(t["balance"]), "right")
        pdf.text(555, y, "1234")
        y += len(lines) * lh
    pdf.text(145, y + 4, "TRANSACTION TOTAL")
    pdf.text(145, y + 14, "CLOSING BALANCE")
    pdf.save()


# ------------------------------------------------------------------ tests

class LayoutTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def check(self, stmt, txns, opening, ref=True, value_date=False):
        got = stmt.transactions
        self.assertEqual(len(got), len(txns), [t.description for t in got][:5])
        for exp, t in zip(txns, got):
            with self.subTest(date=exp["date"], desc=exp["desc"][:20]):
                self.assertEqual(t.date, exp["date"])
                self.assertEqual(t.description, exp.get("shown", exp["desc"]))
                self.assertEqual(t.debit, exp["debit"])
                self.assertEqual(t.credit, exp["credit"])
                self.assertEqual(t.balance, exp["balance"])
                self.assertEqual(t.check, "OK")
                if ref and exp["ref"]:
                    self.assertIn(exp["ref"], t.ref.replace(" ", ""))
                if value_date:
                    self.assertEqual(t.value_date, exp["value_date"])
        self.assertAlmostEqual(stmt.opening_balance, opening, places=2)
        self.assertEqual(stmt.closing_balance, txns[-1]["balance"])
        self.assertEqual(stmt.mismatches, 0)

    def test_hdfc_like_multi_page(self):
        txns = make_transactions(70, seed=2)
        hdfc_like(self.path("h.pdf"), txns, 85000.0)
        stmt = parse_statement(self.path("h.pdf"))
        self.assertEqual(stmt.bank, "HDFC Bank")
        self.assertEqual(stmt.account_no, "50100123456789")
        self.assertEqual(stmt.period, "01-04-2024 to 30-06-2024")
        self.check(stmt, txns, 85000.0, value_date=True)

    def test_sbi_like_ruled_two_line_header(self):
        txns = make_transactions(60, opening=12500.5, seed=3)
        sbi_like(self.path("s.pdf"), txns, 12500.5)
        stmt = parse_statement(self.path("s.pdf"))
        self.assertEqual(stmt.bank, "State Bank of India")
        self.check(stmt, txns, 12500.5, value_date=True)

    def test_kotak_like_amount_with_suffix_and_centred_dates(self):
        txns = make_transactions(50, seed=4)
        kotak_like(self.path("k.pdf"), txns, 85000.0)
        stmt = parse_statement(self.path("k.pdf"))
        self.assertEqual(stmt.bank, "Kotak Mahindra Bank")
        self.check(stmt, txns, 85000.0)
        self.assertEqual(stmt.transactions[0].extra, {"Sl. No.": "1"})

    def test_icici_like_newest_first_header_once(self):
        txns = make_transactions(80, seed=5)
        icici_like(self.path("i.pdf"), txns, 85000.0)
        stmt = parse_statement(self.path("i.pdf"))
        self.assertEqual(stmt.bank, "ICICI Bank")
        self.check(stmt, txns, 85000.0, ref=False)
        self.assertTrue(any("newest-first" in w for w in stmt.warnings))

    def test_icici_like_header_every_page(self):
        txns = make_transactions(80, seed=6)
        icici_like(self.path("i2.pdf"), txns, 85000.0, header_every_page=True)
        self.check(parse_statement(self.path("i2.pdf")), txns, 85000.0, ref=False)

    def test_amount_with_drcr_column(self):
        txns = make_transactions(30, seed=7)
        drcr_column_like(self.path("a.pdf"), txns, 85000.0)
        stmt = parse_statement(self.path("a.pdf"))
        self.assertEqual(stmt.bank, "Axis Bank")
        self.assertEqual(stmt.account_no, "917010012345678")
        self.check(stmt, txns, 85000.0, ref=False)
        self.assertEqual(stmt.extra_columns, ["Init. Br"])

    def test_amount_without_flag_uses_balance(self):
        txns = make_transactions(30, seed=8)
        drcr_column_like(self.path("b.pdf"), txns, 85000.0, with_flag=False)
        self.check(parse_statement(self.path("b.pdf")), txns, 85000.0, ref=False)

    def test_password(self):
        txns = make_transactions(10, seed=9)
        path = self.path("locked.pdf")
        hdfc_like(path, txns, 85000.0, password="ABCD1234")
        with self.assertRaisesRegex(StatementError, "password"):
            parse_statement(path)
        with self.assertRaisesRegex(StatementError, "password"):
            parse_statement(path, password="wrong")
        self.check(parse_statement(path, password="ABCD1234"), txns, 85000.0)

    def test_scanned_pdf_error(self):
        path = self.path("blank.pdf")
        c = canvas.Canvas(path)
        c.rect(50, 50, 100, 100, fill=1)
        c.save()
        with self.assertRaisesRegex(StatementError, "scanned"):
            parse_statement(path)

    def test_excel_output(self):
        txns = make_transactions(20, seed=10)
        hdfc_like(self.path("h.pdf"), txns, 85000.0)
        drcr_column_like(self.path("a.pdf"), txns, 85000.0)
        stmts = [parse_statement(self.path("h.pdf")), parse_statement(self.path("a.pdf"))]
        out = self.path("out.xlsx")
        write_workbook(stmts, out)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Summary", "50100123456789", "917010012345678"])
        ws = wb["50100123456789"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ["Date", "Value Date", "Description / Narration", "Chq. / Ref. No.",
                                  "Debit (Dr)", "Credit (Cr)", "Balance", "Balance Check"])
        self.assertEqual(ws["C2"].value, "Opening Balance")
        self.assertEqual(ws["G2"].value, 85000.0)
        self.assertEqual(ws["A3"].value.date(), txns[0]["date"])
        self.assertEqual(ws["C3"].value, txns[0]["desc"])
        self.assertEqual(ws["H3"].value, "OK")
        last = 2 + len(txns) + 1
        self.assertEqual(ws[f"C{last}"].value, "Total")
        self.assertEqual(ws[f"E{last}"].value, f"=SUM(E3:E{last - 1})")
        summary = wb["Summary"]
        self.assertEqual(summary["C2"].value, "HDFC Bank")
        self.assertEqual(summary["I2"].value, txns[-1]["balance"])
        self.assertEqual(summary["J2"].value, txns[-1]["balance"])
        self.assertEqual(summary["L2"].value, 0)


class ValueParsingTest(unittest.TestCase):
    def test_dates(self):
        cases = {
            "01/04/24": dt.date(2024, 4, 1), "01/04/2024": dt.date(2024, 4, 1),
            "1-4-2024": dt.date(2024, 4, 1), "01.04.2024": dt.date(2024, 4, 1),
            "1 Apr 2024": dt.date(2024, 4, 1), "01-Apr-2024": dt.date(2024, 4, 1),
            "01-APR-24": dt.date(2024, 4, 1), "2024-04-01": dt.date(2024, 4, 1),
            "Apr 1, 2024": dt.date(2024, 4, 1), "01/04/2024 10:32:11": dt.date(2024, 4, 1),
            "Sept 30, 2024": dt.date(2024, 9, 30),
        }
        for text, want in cases.items():
            self.assertEqual(parse_date(text), want, text)
        for text in ("UPI/123/456", "32/01/2024", "Opening", "", "12345678"):
            self.assertIsNone(parse_date(text), text)

    def test_amounts(self):
        cases = {
            "1,23,456.70": (123456.7, None), "500.00": (500.0, None), "(500.00)": (-500.0, None),
            "-12.50": (-12.5, None), "9,000.00 Cr": (9000.0, "CR"), "9,000.00(Dr)": (9000.0, "DR"),
            "₹ 1,000.00": (1000.0, None), "Dr 45.00": (45.0, "DR"), "-": (None, None), "": (None, None),
        }
        for text, want in cases.items():
            self.assertEqual(parse_amount(text), want, text)


if __name__ == "__main__":
    unittest.main()
